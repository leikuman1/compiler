from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "tmp" / "docs" / "实验报告_working.docx"
OUTPUT = ROOT / "tmp" / "docs" / "实验报告_filled.docx"


SECTION_CONTENTS: dict[str, list[str]] = {
    "实验目的": [
        "1. 理解正规式、NFA、DFA 和 MFA 之间的关系，掌握从形式化描述到自动机构造的完整实现过程。",
        "2. 掌握正规式中缀转后缀、Thompson 构造法、子集构造法以及 DFA 最小化算法的核心思想。",
        "3. 通过图形界面、文件读写和状态展示功能，将自动机构造算法落实为可验证、可操作的实验程序。",
    ],
    "二、实验内容": [
        "本次为第二次实验，围绕正规式、NFA、DFA 和 MFA 的构造与转换展开。实验完成了正规式合法性验证、根据正规式生成 NFA、根据 NFA 构造 DFA、在 DFA 最小化前去除冗余状态并生成 MFA，同时实现了 NFA/DFA 文件的读入与保存，以及图形界面的状态转移展示功能。",
    ],
    "三、实验需求（需要实现哪些功能）": [
        "1. 输入正规式并完成语法合法性检查，给出验证成功或失败的提示信息。",
        "2. 将合法正规式转换为后缀表示，并基于 Thompson 构造法生成 NFA。",
        "3. 展示 NFA 的开始状态、终止状态和状态转移表，并支持保存为 .nfa 文件。",
        "4. 将 NFA 通过子集构造法转换为 DFA，并支持 DFA 的读入、展示和保存为 .dfa 文件。",
        "5. 在 DFA 转 MFA 前先去除不可达状态和无效状态，再使用最小化算法得到 MFA。",
        "6. 通过图形界面完成正规式输入、自动机构造、结果查看和文件操作等实验流程。",
    ],
    "四、主要数据结构介绍": [
        "1. RegexAtom：表示正规式中的基本原子，支持普通字符、转义字符以及字符类等输入形式。",
        "2. _Fragment：Thompson 构造过程中的中间片段结构，只保留 start 和 accept 两个关键端点，便于执行连接、并运算和闭包运算。",
        "3. NFA：使用 states、start_state、accept_states、alphabet 和 transitions 记录非确定有限自动机，其中 transitions[source][symbol] = {targets} 表示一对多转移。",
        "4. DFA：使用与 NFA 类似的字段记录确定有限自动机，其中 transitions[source][symbol] = target 表示一对一转移。",
        "5. state_subsets：记录每个 DFA 或 MFA 状态对应的原始状态集合，用于跟踪子集构造与最小化过程。",
        "6. accept_metadata：为接受状态保存附加语义信息，保证状态合并和接受态选择时不丢失原有语义。",
    ],
    "五、主要模块算法介绍": [
        "1. 正规式解析模块 RegexParser：先对输入串进行分词，再自动补全显式连接符，最后依据运算符优先级将中缀正规式转换为后缀表示，为后续栈式构造做准备。",
        "2. 正规式转 NFA 模块 ThompsonBuilder：扫描后缀正规式，对普通字符构造基础片段；对连接、并运算、闭包、正闭包和可选运算分别添加 epsilon 转移，最终拼接成完整 NFA。",
        "3. NFA 转 DFA 模块 Determinizer：以 NFA 初态的 epsilon-closure 作为 DFA 初态，结合 move 与 epsilon-closure 运算扩展新状态，按子集构造法生成确定有限自动机。",
        "4. DFA 转 MFA 模块 DFAMinimizer：首先去除冗余状态，包括从开始状态不可达的状态，以及虽然可达但永远无法到达接受状态的无效状态；随后按接受态和非接受态进行初始划分，并根据各输入符号下的目标分组签名不断细化分组，直到划分稳定，最终得到 MFA。",
        "5. 文件编解码模块 AutomataCodec：负责课程实验格式的 NFA/DFA 文件读写，便于自动机在程序内外部之间交换。",
        "6. 图形界面模块 ui.app：基于 Tkinter 实现正规式验证、自动机构造、状态转移展示和文件操作等交互功能。",
    ],
    "六、程序实现环境及使用说明": [
        "程序实现环境：Windows 平台，Python 3.12.1，图形界面采用 Tkinter，开发与调试可在 PyCharm 等 Python IDE 中完成。",
        "使用说明：启动界面程序后，进入 NFA_DFA_MFA 模块，在输入框中填写正规式并先进行验证；验证通过后依次点击“生成NFA”“生成DFA”“生成MFA”即可查看三类自动机的状态转移结果；NFA 与 DFA 支持按课程实验格式读入和保存，便于后续测试与对比。",
    ],
    "七、实验测试用例设计说明": [
        "1. 基本连接用例：ab。用于验证普通字符连接时的正规式解析、NFA 构造以及 DFA/MFA 转换结果是否正确。",
        "2. 并运算与闭包用例：(a|b)*abb。用于验证括号、并运算、Kleene 闭包和多步转换后的状态集合变化。",
        "3. 正闭包与可选用例：a+b?。用于验证 + 和 ? 运算在 Thompson 构造中的实现是否正确。",
        "4. 字符类用例：[a-z][a-z0-9]*。用于验证字符类、范围表达式以及标识符样式正规式的处理能力。",
        "5. 非法正规式用例：(ab、[a-z、a\\。用于验证括号不匹配、字符类不完整和悬空转义时的报错提示是否正确。",
        "6. 文件测试用例：将生成的 NFA 保存为 .nfa，将生成的 DFA 保存为 .dfa，再重新读入并比对状态转移表，验证文件编解码功能。",
        "7. 冗余状态测试用例：构造或读入带有不可达状态和无效状态的 DFA，检查在 DFA 转 MFA 前是否先完成冗余状态删除，再进行最小化。",
    ],
    "九、个人（或小组）对实验结果的自我评价": [
        "本次第二次实验已经完成正规式验证、正规式转 NFA、NFA 转 DFA 以及 DFA 去冗余后转 MFA 等核心功能，界面交互、文件读写和状态展示也能够正常运行。通过测试，主要算法实现正确，模块划分清晰，能够较好地体现编译原理中自动机构造的基本思想。后续仍可继续补充更丰富的极端用例与展示材料，但从实验要求来看，核心任务已经完成。",
    ],
}


def set_run_font(run, *, name: str = "宋体", size: int = 12, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph: Paragraph, *, first_line_indent: bool = True) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    fmt.first_line_indent = Cm(0.74) if first_line_indent else Cm(0)


def replace_paragraph_text(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def insert_paragraph_after(paragraph: Paragraph, text: str, *, style_name: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style_name
    replace_paragraph_text(new_para, text)
    format_paragraph(new_para)
    return new_para


def style_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_font(run)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def fill_table(doc: Document) -> None:
    table = doc.tables[0]
    table.cell(2, 1).text = "实验二：正规式、NFA、DFA 和 MFA 的构造与转换"
    table.cell(8, 1).text = "2026年4月24日 - 2026年5月8日"
    for row in table.rows:
        for cell in row.cells:
            style_cell(cell)


def fill_section(doc: Document, heading_text: str, contents: list[str]) -> None:
    paragraphs = doc.paragraphs
    heading_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading_text)
    placeholder = paragraphs[heading_index + 1]
    placeholder.style = "Normal"
    replace_paragraph_text(placeholder, contents[0])
    format_paragraph(placeholder)
    current = placeholder
    for text in contents[1:]:
        current = insert_paragraph_after(current, text)


def main() -> None:
    doc = Document(SOURCE)
    fill_table(doc)
    for heading_text, contents in SECTION_CONTENTS.items():
        fill_section(doc, heading_text, contents)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
