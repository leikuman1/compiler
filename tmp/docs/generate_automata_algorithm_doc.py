from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "output" / "doc" / "自动机构造算法说明.docx"


def set_run_font(run, *, name: str = "Microsoft YaHei", size: int | None = None, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    for style_name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    run = paragraph.add_run(text)
    set_run_font(run, size=18, bold=True)
    paragraph.alignment = 1


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.35
    run = paragraph.add_run(text)
    set_run_font(run)


def add_code_block(document: Document, text: str) -> None:
    for line in text.strip("\n").splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=10)


def add_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("正规式构造", "Thompson 片段规则", "实现位置")
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    rows = [
        ("普通字符或字符集", "新建起始状态和终止状态，按字符连一条边", "core/thompson.py::_build_atom"),
        ("连接 AB", "A.accept 通过 epsilon 连到 B.start", "token == CONCAT"),
        ("并运算 A|B", "新建总起点和总终点，分别连向 A 和 B", "token == '|'"),
        ("闭包 A*", "新建总起点和总终点，加入回边和空串通路", "token == '*'"),
        ("正闭包 A+", "至少执行一次，再从 accept 回到 start", "token == '+'"),
        ("可选 A?", "允许走 A，也允许直接走空串", "token == '?'"),
    ]

    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10)


def build_document() -> Document:
    document = Document()
    style_document(document)

    add_title(document, "正规式生成 NFA、NFA 转 DFA、DFA 转 MFA 算法说明")

    add_paragraph(
        document,
        "本文档结合当前项目代码，系统讲解三个核心算法：根据正规式构造 NFA、由 NFA 构造 DFA，以及由 DFA 进一步最小化得到 MFA。文档不仅说明算法原理，也解释项目中的具体数据结构、状态表示方式和关键实现细节，便于课程报告、代码阅读和功能答辩使用。",
    )

    document.add_heading("1. 项目中的自动机数据结构", level=1)
    add_paragraph(
        document,
        "本项目在 core/models.py 中定义了 NFA 和 DFA 两种数据结构。NFA 的转移表形式为 transitions[source][symbol] = {target1, target2, ...}，因此同一个状态在同一个输入符号下可以到达多个后继状态；DFA 的转移表形式为 transitions[source][symbol] = target，因此每个输入符号最多只有一个后继状态。项目把 epsilon 边记为字符 #，这意味着空串转移不会出现在字母表里，但会参与闭包计算与 Thompson 构造。",
    )
    add_bullet(document, "NFA 重点字段：states、start_state、accept_states、alphabet、transitions、accept_metadata。")
    add_bullet(document, "DFA 重点字段：在 NFA 基础上额外保留 state_subsets，用来记录每个 DFA 状态对应的 NFA 状态集合。")
    add_bullet(document, "accept_metadata 用于词法分析器场景下的优先级与类别保留，因此 determinize 和 minimize 都考虑了接受状态的附加信息。")

    document.add_heading("2. 正规式预处理：从中缀表达式到后缀表示", level=1)
    add_paragraph(
        document,
        "在真正构造 NFA 之前，项目先通过 core/regex.py 中的 RegexParser 对输入正规式进行预处理。这个步骤包括分词、补全显式连接符和转换为后缀表达式三部分。后缀表达式的好处是可以用一个栈按顺序处理操作符，天然适合 Thompson 构造。",
    )
    add_bullet(document, "分词：识别普通字符、转义字符、字符类 [a-z]、括号、并运算 | 和一元运算符 *、+、?。")
    add_bullet(document, "插入连接符：例如 ab 会被显式改写成 a.b，(a|b)c 会被改写成 (a|b).c。")
    add_bullet(document, "中缀转后缀：使用类似 Shunting Yard 的思想，根据优先级把表达式改写成便于栈处理的后缀序列。")
    add_paragraph(
        document,
        "以正规式 (a|b)*abb 为例，项目会先补出连接符，得到 (a|b)*.a.b.b，然后再转换为后缀形式 a b | * a . b . b .。后续 ThompsonBuilder 正是按这个后缀序列逐项处理。",
    )

    document.add_heading("3. 根据正规式生成 NFA：Thompson 构造法", level=1)
    add_paragraph(
        document,
        "core/thompson.py 中的 ThompsonBuilder 采用 Thompson 构造法。该算法的核心思想是：为每一个基本字符构造一个最小片段，然后在扫描后缀正规式时，用栈不断把小片段组合成大片段。每个片段都只有两个暴露端点，即 start 和 accept，这种结构非常适合做连接、并运算和闭包运算。",
    )
    add_paragraph(
        document,
        "项目中用 _Fragment(start, accept) 表示一个中间自动机片段。读到普通字符时，调用 _build_atom 新建两个状态，并在它们之间建立带符号的转移；读到操作符时，则弹出一个或两个已有片段，再按运算类型加入 epsilon 转移，组合成一个新的片段压回栈中。所有 token 处理结束后，栈中应只剩下一个片段，它的起点就是 NFA 的开始状态，它的终点就是 NFA 的接受状态。",
    )
    add_table(document)
    add_paragraph(
        document,
        "这个实现的一个特点是所有状态编号都由 _new_state 顺序生成，因此状态编号稳定、便于调试和导出图形。另外，_add_transition 会自动维护状态集合和字母表集合，使得最终返回的 NFA 结构已经是完整可用的自动机对象。",
    )
    add_paragraph(document, "Thompson 构造过程可以概括为下面的伪代码：")
    add_code_block(
        document,
        """
for token in postfix_regex:
    if token is atom:
        push(atom_fragment(token))
    elif token is CONCAT:
        right = pop()
        left = pop()
        connect(left.accept, epsilon, right.start)
        push(Fragment(left.start, right.accept))
    elif token is UNION:
        create new start and accept
        connect(new_start, epsilon, left.start and right.start)
        connect(left.accept and right.accept, epsilon, new_accept)
        push(Fragment(new_start, new_accept))
    elif token is closure operator:
        create new structure according to *, + or ?
        push(new_fragment)
result = stack only fragment
        """,
    )

    document.add_heading("4. NFA 转 DFA：子集构造法", level=1)
    add_paragraph(
        document,
        "core/determinize.py 中的 Determinizer 使用子集构造法把 epsilon-NFA 转换成 DFA。该算法的基本思想是：DFA 中的一个状态，不再对应原 NFA 的单个状态，而是对应 NFA 状态的一个集合。只要集合表示确定，那么在任意输入符号下的转移也就确定，从而可以得到确定有穷自动机。",
    )
    add_bullet(document, "epsilon_closure(states)：求一个状态集合在任意多条 epsilon 边作用下能够到达的全部状态。")
    add_bullet(document, "move(states, symbol)：求状态集合在读入某个输入符号后能到达的所有状态。")
    add_bullet(document, "subset_to_state：把 NFA 状态子集映射为新的 DFA 状态编号，避免重复创建。")
    add_bullet(document, "queue：使用广度优先方式逐步展开新的 DFA 状态。")
    add_paragraph(
        document,
        "算法从 NFA 起始状态的 epsilon 闭包开始，把它作为 DFA 的 0 号状态。然后对队列中的每一个子集，枚举字母表中的每一个输入符号，先做 move，再做 epsilon 闭包，得到目标子集。如果该目标子集以前没有出现过，就分配一个新的 DFA 状态编号并入队；如果出现过，则直接复用旧编号。不断重复这个过程，直到队列为空，整个 DFA 便构造完成。",
    )
    add_paragraph(
        document,
        "本项目还处理了接受状态元数据的问题。因为多个 NFA 接受状态在子集构造后可能会合并到同一个 DFA 状态里，所以 _best_accept 会根据 accept_choice_key 选择一个最优的接受动作。这一点对词法分析尤其重要，因为不同 token 的优先级和先后顺序可能不同。",
    )
    add_paragraph(document, "子集构造法的伪代码可以写成：")
    add_code_block(
        document,
        """
start_subset = epsilon_closure({nfa.start_state})
assign DFA state 0 to start_subset
push start_subset into queue

while queue not empty:
    subset = pop_left(queue)
    for symbol in alphabet:
        moved = move(subset, symbol)
        if moved is empty:
            continue
        target_subset = epsilon_closure(moved)
        if target_subset is new:
            assign new DFA state id
            push target_subset into queue
        add DFA transition(current_state, symbol, target_state)
        """,
    )

    document.add_heading("5. DFA 转 MFA：迭代划分最小化", level=1)
    add_paragraph(
        document,
        "core/minimize.py 中的 DFAMinimizer 使用迭代划分 refinement 的方法最小化 DFA。其思想是：如果两个状态在所有输入符号下的行为完全等价，那么它们可以合并为一个状态；如果存在某个输入符号会把它们带到不同类别的状态中，那么它们就不能合并。",
    )
    add_paragraph(
        document,
        "算法首先构造初始划分。项目不是简单地把状态分成接受状态和非接受状态两类，而是进一步按 accept_metadata 对接受状态分组。这样可以保证词法分析器中语义不同的接受状态不会被错误合并。完成初始划分后，算法不断检查每个分组内部的状态：比较它们在每个输入符号下所到达的目标分组编号，并把目标分组编号序列作为签名 signature。签名相同的状态保持在同一组，签名不同的状态被拆到不同组。只要某一轮发生了拆分，就继续下一轮，直到划分稳定为止。",
    )
    add_bullet(document, "初始划分：接受状态按 metadata 分组，非接受状态单独成组。")
    add_bullet(document, "划分细化：按每个状态在各输入符号下的目标分组编号生成 signature。")
    add_bullet(document, "稳定终止：某一轮没有任何分裂时，说明已经无法继续区分。")
    add_bullet(document, "重建 MFA：每个最终分组对应一个新状态，用组内代表状态的转移来构造新自动机。")
    add_paragraph(
        document,
        "项目还调用 _order_partitions 把包含开始状态的分组排在最前面，这样最小化后的开始状态总是编号 0，对输出结果和界面展示都更友好。同时 state_subsets 字段会记录每个新状态对应的原 DFA 状态集合，便于回溯合并关系。",
    )
    add_paragraph(document, "最小化过程的伪代码如下：")
    add_code_block(
        document,
        """
partitions = initial_partitions(dfa)
changed = True

while changed:
    changed = False
    new_partitions = []
    for partition in partitions:
        split partition by signature of outgoing transitions
        if partition is split:
            changed = True
        append all buckets to new_partitions
    partitions = new_partitions

rebuild minimized DFA from final partitions
        """,
    )

    document.add_heading("6. 三个算法之间的关系", level=1)
    add_paragraph(
        document,
        "这三个算法构成了一个逐步规范化的流程。首先，正规式经过解析与 Thompson 构造，被转换为一个容易生成但通常含有大量 epsilon 边的 NFA；接着，子集构造法把 NFA 转换为等价的 DFA，使得每一步输入都有唯一后继，方便程序执行；最后，DFA 最小化进一步合并等价状态，减少状态数量，得到结构更紧凑、运行效率更高的 MFA。三步的目标并不相同：第一步强调可构造性，第二步强调确定性，第三步强调最简性。",
    )

    document.add_heading("7. 与项目代码文件的对应关系", level=1)
    add_bullet(document, "正规式解析：core/regex.py")
    add_bullet(document, "Thompson 构造 NFA：core/thompson.py")
    add_bullet(document, "子集构造 NFA -> DFA：core/determinize.py")
    add_bullet(document, "DFA 最小化：core/minimize.py")
    add_bullet(document, "自动机数据结构：core/models.py")
    add_bullet(document, "界面触发流程：ui/app.py 中的 generate_nfa、generate_dfa、generate_mfa")

    document.add_heading("8. 总结", level=1)
    add_paragraph(
        document,
        "从课程实验的角度看，这套实现结构清晰、职责划分明确：RegexParser 负责把输入正规式变成适合处理的后缀表示，ThompsonBuilder 负责构造 NFA，Determinizer 负责完成确定化，DFAMinimizer 负责进一步压缩状态数量。由于每一步都保留了足够的数据结构信息，这个项目不仅能完成自动机构造，还能支持词法分析、自动机导出和界面展示等上层功能。",
    )

    return document


def add_metadata(document: Document) -> None:
    props = document.core_properties
    props.title = "正规式生成 NFA、NFA 转 DFA、DFA 转 MFA 算法说明"
    props.author = "OpenAI Codex"
    props.subject = "Compiler Project Automata Algorithms"
    props.comments = "Generated from the current project implementation."


def main() -> None:
    document = build_document()
    add_metadata(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
